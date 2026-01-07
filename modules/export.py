from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from datetime import datetime
from typing import List, Dict, Any


class ExportManager:
    """데이터 내보내기 관리 클래스"""
    
    def __init__(self):
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.functional_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def export_to_excel(self, test_items: List[Dict[str, Any]], user_name: str, project_name: str) -> bytes:
        """
        시험 항목을 Excel 파일로 내보내기
        
        Args:
            test_items: 시험 항목 리스트
            user_name: 사용자 이름
            project_name: 프로젝트 이름
            
        Returns:
            Excel 파일 바이트
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Plan"
        
        # 헤더 정의
        headers = [
            "No", "Group", "Specification", "§", "Test Title", "Component",
            "Test by.", "Sample quantity", "Test Timing", "Start", "End",
            "OK/NOK", "Result", "Remark", "Customer Feedback"
        ]
        
        # 헤더 작성
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.border
        
        # 그룹별 분류
        grouped_data = self._group_test_items(test_items)
        
        # 데이터 작성
        row_num = 2
        item_no = 1
        
        for category_idx, (category, items) in enumerate(grouped_data.items(), 1):
            # 섹션 헤더
            section_cell = ws.cell(row=row_num, column=1)
            section_cell.value = category
            section_cell.font = Font(bold=True, size=11)
            section_cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
            ws.merge_cells(f'A{row_num}:O{row_num}')
            row_num += 1
            
            # 항목 작성
            for item_idx, item in enumerate(items, 1):
                # Functional Test 여부 확인
                is_functional = 'functional' in item.get('test_name', '').lower()
                
                # No
                ws.cell(row=row_num, column=1).value = item_no
                
                # Group (Sequence/Individual 판단)
                group = self._determine_group(item, items)
                ws.cell(row=row_num, column=2).value = group
                
                # Specification
                ws.cell(row=row_num, column=3).value = item.get('ref_standard', '')
                
                # § (자동 넘버링)
                section_num = f"{category_idx}.{item_idx}.1"
                ws.cell(row=row_num, column=4).value = section_num
                
                # Test Title
                ws.cell(row=row_num, column=5).value = item.get('test_name', '')
                
                # Component
                ws.cell(row=row_num, column=6).value = item.get('sample_assembly', '')
                
                # Test by.
                ws.cell(row=row_num, column=7).value = user_name
                
                # Sample quantity
                ws.cell(row=row_num, column=8).value = item.get('sample_count', '3')
                
                # Test Timing
                ws.cell(row=row_num, column=9).value = ""
                
                # Start, End
                ws.cell(row=row_num, column=10).value = ""
                ws.cell(row=row_num, column=11).value = ""
                
                # OK/NOK, Result, Remark, Customer Feedback
                for col in range(12, 16):
                    ws.cell(row=row_num, column=col).value = ""
                
                # 스타일 적용
                for col in range(1, 16):
                    cell = ws.cell(row=row_num, column=col)
                    cell.border = self.border
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    
                    # Functional Test 하이라이트
                    if is_functional:
                        cell.fill = self.functional_fill
                
                row_num += 1
                item_no += 1
        
        # 시료 수 계산 및 경고
        total_samples = self._calculate_total_samples(test_items)
        warning_row = row_num + 1
        ws.cell(row=warning_row, column=1).value = f"총 시료 수: {total_samples}"
        ws.cell(row=warning_row, column=1).font = Font(bold=True, color="FF0000")
        
        # 열 너비 조정
        column_widths = [5, 12, 15, 8, 30, 15, 15, 12, 12, 10, 10, 10, 15, 20, 20]
        for idx, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        
        # 파일 저장
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue()
    
    def _group_test_items(self, test_items: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
        """시험 항목을 카테고리별로 그룹화"""
        grouped = {}
        
        for item in test_items:
            category = item.get('category', 'Other')
            if category not in grouped:
                grouped[category] = []
            grouped[category].append(item)
        
        # Functional Test를 최상단으로
        if 'Functional Test' in grouped:
            functional = grouped.pop('Functional Test')
            grouped = {'Functional Test': functional, **grouped}
        
        return grouped
    
    def _determine_group(self, item: Dict[str, Any], all_items: List[Dict[str, Any]]) -> str:
        """시험 항목의 그룹 결정 (Sequence/Individual)"""
        sample_no = item.get('test_sample_no', '')
        
        if not sample_no:
            return 'Individual'
        
        # 같은 샘플 번호를 가진 항목이 여러 개인지 확인
        same_sample_count = sum(1 for i in all_items if i.get('test_sample_no') == sample_no)
        
        return 'Sequence' if same_sample_count > 1 else 'Individual'
    
    def _calculate_total_samples(self, test_items: List[Dict[str, Any]]) -> int:
        """전체 시료 수 계산"""
        total = 0
        
        # 샘플 번호별로 그룹화
        sample_groups = {}
        for item in test_items:
            sample_no = item.get('test_sample_no', f"individual_{id(item)}")
            if sample_no not in sample_groups:
                sample_groups[sample_no] = []
            sample_groups[sample_no].append(item)
        
        # 각 그룹당 3개씩 계산
        for sample_no, items in sample_groups.items():
            if len(items) > 1:  # Sequence
                total += 3
            else:  # Individual
                total += 3
        
        return total

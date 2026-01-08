from typing import Dict, List
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

class PlanningManager:
    def __init__(self):
        self.section_counters = {}
    
    def create_draft_plan(self, test_items: List[Dict]) -> pd.DataFrame:
        """계획서 초안 생성"""
        # 그룹 판별
        grouped_items = self._group_test_items(test_items)
        
        # Functional Test 확인 및 생성
        has_functional = any(item.get('test_name', '').lower() == 'functional test' 
                            for item in test_items)
        
        if not has_functional:
            functional_item = self._generate_functional_test(test_items)
            grouped_items.insert(0, functional_item)
        
        # 계획서 테이블 생성
        plan_rows = []
        self.section_counters = {}
        
        for idx, item in enumerate(grouped_items, 1):
            category = item.get('category', 'Other')
            section_num = self._get_section_number(category)
            
            row = {
                'No': idx,
                'Group': item.get('group_type', 'Individual'),
                'Specification': self._generate_spec_name(category, section_num),
                '§': section_num,
                'Test Title': item.get('test_name', ''),
                'Component': item.get('sample_assembly', ''),
                'Test by.': '',
                'Sample quantity': item.get('sample_count', '3'),
                'Test Timing': '',
                'Start': '',
                'End': '',
                'OK/NOK': '',
                'Result': '',
                'Remark': '',
                'Customer Feedback': ''
            }
            plan_rows.append(row)
        
        return pd.DataFrame(plan_rows)
    
    def _group_test_items(self, test_items: List[Dict]) -> List[Dict]:
        """시험 항목 그룹화"""
        sample_groups = {}
        
        for item in test_items:
            sample_no = item.get('test_sample_no', '')
            if sample_no:
                if sample_no not in sample_groups:
                    sample_groups[sample_no] = []
                sample_groups[sample_no].append(item)
        
        result = []
        for sample_no, items in sample_groups.items():
            if len(items) > 1:
                for item in items:
                    item['group_type'] = 'Sequence'
            else:
                items[0]['group_type'] = 'Individual'
            result.extend(items)
        
        # 샘플 번호 없는 항목들
        for item in test_items:
            if not item.get('test_sample_no'):
                item['group_type'] = 'Individual'
                result.append(item)
        
        return result
    
    def _generate_functional_test(self, test_items: List[Dict]) -> Dict:
        """Functional Test 자동 생성"""
        # 가장 많이 출현한 컴포넌트 찾기
        components = [item.get('sample_assembly', '') for item in test_items if item.get('sample_assembly')]
        most_common = max(set(components), key=components.count) if components else 'Motor only'
        
        # 전체 시료 수 계산
        total_samples = sum(int(item.get('sample_count', 3)) for item in test_items)
        
        return {
            'test_name': 'Functional Test',
            'category': 'Operational and Environmental tests',
            'sample_assembly': most_common,
            'sample_count': str(total_samples),
            'test_duration': '1',
            'group_type': 'Individual',
            'test_master_id': 'M006',
            'ref_standard': '',
            'custom_specs': {}
        }
    
    def _get_section_number(self, category: str) -> str:
        """섹션 번호 생성"""
        # 카테고리별 메인 번호
        category_main = {
            'Operational and Environmental tests': '1',
            'Electrical Tests': '2',
            'Endurance Test': '3'
        }
        
        main_num = category_main.get(category, '4')
        
        if category not in self.section_counters:
            self.section_counters[category] = {'sub': 1, 'item': 1}
        else:
            self.section_counters[category]['item'] += 1
        
        return f"{main_num}.{self.section_counters[category]['sub']}.{self.section_counters[category]['item']}"
    
    def _generate_spec_name(self, category: str, section_num: str) -> str:
        """Specification 이름 생성"""
        main_num = section_num.split('.')[0]
        sub_num = section_num.split('.')[1]
        
        if category == 'Electrical Tests':
            return f"Electrical Test spec {main_num}.{sub_num}"
        elif category == 'Operational and Environmental tests':
            return f"Environmental Test spec {main_num}.{sub_num}"
        else:
            return f"Test spec {main_num}.{sub_num}"
    
    def export_to_excel(self, df: pd.DataFrame, user_name: str) -> str:
        """엑셀 파일로 내보내기"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Test_Plan_{user_name}_{timestamp}.xlsx"
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Plan"
        
        # 헤더 작성
        headers = df.columns.tolist()
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 데이터 작성
        for row_idx, row in enumerate(df.itertuples(index=False), 2):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Functional Test 하이라이트
                if col_idx == 5 and value == 'Functional Test':
                    for c in range(1, len(headers) + 1):
                        ws.cell(row=row_idx, column=c).fill = PatternFill(
                            start_color="FFFF00", end_color="FFFF00", fill_type="solid"
                        )
        
        # 열 너비 조정
        column_widths = {
            'A': 5, 'B': 12, 'C': 25, 'D': 10, 'E': 30,
            'F': 15, 'G': 12, 'H': 12, 'I': 12, 'J': 12,
            'K': 12, 'L': 10, 'M': 15, 'N': 20, 'O': 20
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        wb.save(filename)
        return filename

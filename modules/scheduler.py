import plotly.figure_factory as ff
import plotly.graph_objects as go
import pandas as pd
from datetime import datetime, timedelta
from io import BytesIO

class Scheduler:
    def __init__(self, db):
        self.db = db
        
        # 우선순위 정의 (낮을수록 먼저)
        self.category_priority = {
            "Electrical Tests": 1,
            "Operational and Environmental Tests": 2,
            "Endurance Test": 3,
            "Performance Test": 4,
        }
    
    def sort_by_priority(self, test_items):
        """시험 항목을 우선순위에 따라 정렬"""
        def get_priority(item):
            category = item.get('category', '')
            return self.category_priority.get(category, 999)
        
        return sorted(test_items, key=get_priority)
    
    def create_monthly_gantt(self, schedule_data):
        """월별 Gantt 차트 생성"""
        if not schedule_data:
            return go.Figure()
        
        df_data = []
        for item in schedule_data:
            df_data.append({
                'Task': item['test_name'],
                'Start': item['start_date'],
                'Finish': item['end_date'],
                'Resource': item.get('category', 'Unknown')
            })
        
        df = pd.DataFrame(df_data)
        
        fig = ff.create_gantt(
            df,
            colors=['#FF6B6B', '#4ECDC4', '#45B7D1', '#FFA07A'],
            index_col='Resource',
            show_colorbar=True,
            group_tasks=True,
            showgrid_x=True,
            showgrid_y=True
        )
        
        fig.update_layout(
            title="월별 시험 일정",
            xaxis_title="날짜",
            height=400
        )
        
        return fig
    
    def create_request_gantt(self, schedule_data):
        """의뢰별 Gantt 차트 생성"""
        return self.create_monthly_gantt(schedule_data)
    
    def create_dday_gantt(self, df, start_date):
        """D-Day 기준 Gantt 차트 생성"""
        fig_data = []
        
        current_date = start_date
        colors = ['#FF6B6B', '#4ECDC4', '#45B7D1', '#FFA07A', '#96CEB4']
        
        for idx, row in df.iterrows():
            duration = row['소요일수']
            end_date = current_date + timedelta(days=duration)
            
            fig_data.append({
                'Task': row['시험명'],
                'Start': current_date.strftime('%Y-%m-%d'),
                'Finish': end_date.strftime('%Y-%m-%d'),
                'Resource': row['분류']
            })
            
            current_date = end_date
        
        if not fig_data:
            return go.Figure()
        
        df_gantt = pd.DataFrame(fig_data)
        
        fig = ff.create_gantt(
            df_gantt,
            colors=colors,
            index_col='Resource',
            show_colorbar=True,
            group_tasks=True,
            showgrid_x=True,
            showgrid_y=True
        )
        
        fig.update_layout(
            title="시험 일정 타임라인 (D-Day 기준)",
            xaxis_title="날짜",
            yaxis_title="시험 항목",
            height=max(400, len(fig_data) * 30)
        )
        
        return fig
    
    def prepare_schedule_data(self, df, start_date, request_id):
        """일정 데이터 준비"""
        schedule_data = []
        current_date = start_date
        
        for idx, row in df.iterrows():
            duration = row['소요일수']
            end_date = current_date + timedelta(days=duration)
            
            schedule_data.append({
                'test_name': row['시험명'],
                'category': row['분류'],
                'start_date': current_date.isoformat(),
                'end_date': end_date.isoformat(),
                'duration_days': duration,
                'request_id': request_id
            })
            
            current_date = end_date
        
        return schedule_data
    
    def create_excel_plan(self, df, request):
        """Excel 계획서 생성"""
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        
        wb = Workbook()
        ws = wb.active
        ws.title = "시험 계획서"
        
        # 헤더
        ws['A1'] = "시험 계획서"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:K1')
        
        ws['A2'] = f"발주처: {request.get('client', 'N/A')}"
        ws['A3'] = f"프로젝트: {request.get('project', 'N/A')}"
        ws['A4'] = f"작성일: {datetime.now().strftime('%Y-%m-%d')}"
        
        # 테이블 헤더
        headers = list(df.columns)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # 데이터
        for row_idx, row_data in enumerate(df.itertuples(index=False), 7):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # 파일 저장
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue()

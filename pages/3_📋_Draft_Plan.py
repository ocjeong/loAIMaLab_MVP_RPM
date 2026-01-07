import streamlit as st
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO
from modules.data_manager import DataManager

st.set_page_config(page_title="Draft Plan", page_icon="📋", layout="wide")

dm = DataManager()

st.title("📋 계획서 초안 작성")

# 세션 체크
if not st.session_state.current_user:
    st.warning("⚠️ 먼저 사용자를 선택하세요.")
    st.stop()

if not st.session_state.extracted_data:
    st.warning("⚠️ 추출된 데이터가 없습니다.")
    st.stop()

# 데이터 로드
extracted_data = st.session_state.extracted_data
test_items = extracted_data.get('test_items', [])

# 계획서 초안 생성 함수
def create_draft_plan(test_items):
    """계획서 초안 생성"""
    plan_items = []
    
    # Functional Test 자동 생성 확인
    has_functional = any('functional' in item.get('test_name', '').lower() for item in test_items)
    
    # 가장 많이 출현한 컴포넌트 찾기
    components = [item.get('sample_assembly', '') for item in test_items if item.get('sample_assembly')]
    most_common_component = max(set(components), key=components.count) if components else 'Motor only'
    
    # Functional Test 추가
    if not has_functional:
        plan_items.append({
            'No': 1,
            'Group': 'Individual',
            'Specification': 'Functional Test spec 0.1',
            '§': '0.1.1',
            'Test Title': 'Functional Test',
            'Component': most_common_component,
            'Test by': st.session_state.current_user.get('user_name', ''),
            'Sample quantity': 3,
            'Test Timing': 'Before/After each test',
            'Start': '',
            'End': '',
            'OK/NOK': '',
            'Result': '',
            'Remark': 'Auto-generated',
            'Customer Feedback': ''
        })
    
    # 카테고리별 그룹화
    categories = {}
    for item in test_items:
        category = item.get('category', 'Other')
        if category not in categories:
            categories[category] = []
        categories[category].append(item)
    
    # 카테고리별 번호 매핑
    category_numbers = {
        'Environmental Test': 1,
        'Endurance Test': 1,
        'Operational and Environmental tests': 1,
        'Electrical Tests': 2,
        'Performance Test': 3
    }
    
    no = 2 if not has_functional else 1
    
    for category, items in categories.items():
        cat_num = category_numbers.get(category, 4)
        sub_num = 1
        
        for item in items:
            # 그룹 판별 (샘플 번호 기준)
            sample_no = item.get('test_sample_no', '')
            group = 'Sequence' if sample_no else 'Individual'
            
            # § 번호 생성
            section_num = f"{cat_num}.1.{sub_num}"
            
            # Specification 생성
            if category == 'Electrical Tests':
                spec = f"Electrical Test spec {cat_num}.1"
            else:
                spec = f"{category} spec {cat_num}.1"
            
            plan_item = {
                'No': no,
                'Group': group,
                'Specification': spec,
                '§': section_num,
                'Test Title': item.get('test_name', ''),
                'Component': item.get('sample_assembly', ''),
                'Test by': st.session_state.current_user.get('user_name', ''),
                'Sample quantity': item.get('sample_count', 3),
                'Test Timing': '',
                'Start': '',
                'End': '',
                'OK/NOK': '',
                'Result': '',
                'Remark': '',
                'Customer Feedback': ''
            }
            
            plan_items.append(plan_item)
            no += 1
            sub_num += 1
    
    return plan_items

# 계획서 초안 생성
if 'draft_plan' not in st.session_state or st.session_state.draft_plan is None:
    st.session_state.draft_plan = create_draft_plan(test_items)

draft_plan = st.session_state.draft_plan

# 계획서 표시
st.header("📊 시험 계획서 초안")

# 통계 정보
col1, col2, col3, col4 = st.columns(4)
with col1:
    st.metric("총 항목", len(draft_plan))
with col2:
    sequence_count = sum(1 for item in draft_plan if item['Group'] == 'Sequence')
    st.metric("Sequence 그룹", sequence_count)
with col3:
    individual_count = sum(1 for item in draft_plan if item['Group'] == 'Individual')
    st.metric("Individual 항목", individual_count)
with col4:
    total_samples = sum(int(item['Sample quantity']) if isinstance(item['Sample quantity'], (int, str)) and str(item['Sample quantity']).isdigit() else 0 for item in draft_plan)
    st.metric("총 시료 수", total_samples)

st.divider()

# 데이터프레임으로 표시
df = pd.DataFrame(draft_plan)

# 편집 가능한 데이터 에디터
edited_df = st.data_editor(
    df,
    use_container_width=True,
    num_rows="dynamic",
    column_config={
        "No": st.column_config.NumberColumn("No", width="small"),
        "Group": st.column_config.SelectboxColumn(
            "Group",
            options=["Individual", "Sequence"],
            width="small"
        ),
        "Specification": st.column_config.TextColumn("Specification", width="medium"),
        "§": st.column_config.TextColumn("§", width="small"),
        "Test Title": st.column_config.TextColumn("Test Title", width="large"),
        "Component": st.column_config.TextColumn("Component", width="medium"),
        "Test by": st.column_config.TextColumn("Test by", width="small"),
        "Sample quantity": st.column_config.NumberColumn("Sample quantity", width="small"),
        "Test Timing": st.column_config.TextColumn("Test Timing", width="medium"),
        "OK/NOK": st.column_config.SelectboxColumn(
            "OK/NOK",
            options=["", "OK", "NOK"],
            width="small"
        ),
    },
    height=600
)

# 편집된 데이터 저장
st.session_state.draft_plan = edited_df.to_dict('records')

st.divider()

# 엑셀 생성 함수
def create_excel(draft_plan, user_name):
    """엑셀 파일 생성"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Plan"
    
    # 헤더
    headers = list(draft_plan[0].keys())
    ws.append(headers)
    
    # 헤더 스타일
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # 데이터 입력
    for item in draft_plan:
        ws.append(list(item.values()))
    
    # Functional Test 하이라이트
    functional_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if 'functional' in str(row[4].value).lower():  # Test Title 컬럼
            for cell in row:
                cell.fill = functional_fill
    
    # 열 너비 조정
    column_widths = {
        'A': 8, 'B': 12, 'C': 25, 'D': 10, 'E': 30,
        'F': 15, 'G': 12, 'H': 12, 'I': 15, 'J': 12,
        'K': 12, 'L': 10, 'M': 15, 'N': 15, 'O': 20
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # 테두리
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.border = thin_border
            if cell.row > 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # BytesIO로 저장
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    return excel_file

# 버튼
col1, col2, col3 = st.columns([2, 1, 1])

with col1:
    if st.button("🔄 샘플 데이터 로드", use_container_width=True):
        st.session_state.draft_plan = create_draft_plan(test_items)
        st.rerun()

with col2:
    # 엑셀 저장
    user_name = st.session_state.current_user.get('user_name', 'User')
    date_str = datetime.now().strftime('%Y%m%d')
    filename = f"Test_Plan_{user_name}_{date_str}.xlsx"
    
    excel_file = create_excel(edited_df.to_dict('records'), user_name)
    
    st.download_button(
        label="💾 엑셀 저장",
        data=excel_file,
        file_name=filename,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

with col3:
    if st.button("📅 일정 생성", use_container_width=True, type="primary"):
        st.session_state.draft_plan = edited_df.to_dict('records')
        st.success("✅ 계획서가 저장되었습니다!")
        st.info("👉 '📅 Schedule' 페이지로 이동하세요.")

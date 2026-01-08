import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from config import DEFAULT_SAMPLE_COUNT

class PlanGenerator:
    def __init__(self, db_manager):
        self.db_manager = db_manager
    
    def generate_plan(self, test_items):
        """계획서 초안 생성"""
        # Functional Test 자동 생성
        test_items = self._add_functional_test(test_items)
        
        # 그룹 분류
        test_items = self._classify_groups(test_items)
        
        # 계획서 테이블 생성
        plan_table = self._create_plan_table(test_items)
        
        return plan_table
    
    def _add_functional_test(self, test_items):
        """Functional Test 자동 생성"""
        # 이미 Functional Test가 있는지 확인
        has_functional = any(
            item.get('category', '').lower() == 'functional test' or
            'functional' in item.get('test_name', '').lower()
            for item in test_items
        )
        
        if has_functional:
            return test_items
        
        # 가장 많이 출현한 컴포넌트 찾기
        components = [item.get('sample_assembly', '') for item in test_items if item.get('sample_assembly')]
        if components:
            most_common = max(set(components), key=components.count)
        else:
            most_common = 'Motor only'
        
        # 전체 시료 수 계산
        total_samples = sum(int(item.get('sample_count', DEFAULT_SAMPLE_COUNT)) 
                          for item in test_items 
                          if str(item.get('sample_count', '')).isdigit())
        
        # Functional Test 항목 생성
        functional_test = {
            'test_name': 'Functional Test',
            'test_name_original': 'Functional Test',
            'category': 'Functional Test',
            'category_original': 'Functional Test',
            'ref_standard': '',
            'sample_assembly': most_common,
            'test_sample_no': '',
            'sample_count': str(total_samples) if total_samples > 0 else str(DEFAULT_SAMPLE_COUNT),
            'test_duration': '1',
            'test_equipment': '',
            'test_master_id': 'M011',
            'custom_specs': {}
        }
        
        # 맨 앞에 추가
        return [functional_test] + test_items
    
    def _classify_groups(self, test_items):
        """시험 항목 그룹 분류"""
        # 샘플 번호별로 그룹화
        sample_groups = {}
        for item in test_items:
            sample_no = item.get('test_sample_no', '')
            if sample_no:
                if sample_no not in sample_groups:
                    sample_groups[sample_no] = []
                sample_groups[sample_no].append(item)
        
        # 그룹 분류
        for item in test_items:
            sample_no = item.get('test_sample_no', '')
            if sample_no and len(sample_groups.get(sample_no, [])) > 1:
                item['group_type'] = 'Sequence'
            else:
                item['group_type'] = 'Individual'
        
        return test_items
    
    def _create_plan_table(self, test_items):
        """계획서 테이블 생성"""
        plan_data = []
        category_counters = {}
        
        for idx, item in enumerate(test_items):
            category = item.get('category', 'Other')
            
            # 카테고리별 카운터 초기화
            if category not in category_counters:
                category_counters[category] = {'major': len(category_counters) + 1, 'minor': 0}
            
            category_counters[category]['minor'] += 1
            
            # § 번호 생성
            section_num = f"{category_counters[category]['major']}.1.{category_counters[category]['minor']}"
            
            # Specification 생성
            if 'electrical' in category.lower():
                spec = f"Electrical Test spec {category_counters[category]['major']}.1"
            else:
                spec = f"{category} spec {category_counters[category]['major']}.1"
            
            plan_row = {
                'No': idx + 1,
                'Group': item.get('group_type', 'Individual'),
                'Specification': spec,
                '§': section_num,
                'Test Title': item.get('test_name', ''),
                'Component': item.get('sample_assembly', ''),
                'Test by': '',
                'Sample quantity': item.get('sample_count', DEFAULT_SAMPLE_COUNT),
                'Test Timing': '',
                'Start': '',
                'End': '',
                'OK/NOK': '',
                'Result': '',
                'Remark': '',
                'Customer Feedback': '',
                '_category': category,
                '_is_functional': item.get('category', '').lower() == 'functional test',
                '_original_item': item
            }
            
            plan_data.append(plan_row)
        
        return plan_data
    
    def export_to_excel(self, plan_data, user_name):
        """엑셀 파일로 내보내기"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Plan"
        
        # 헤더
        headers = ['No', 'Group', 'Specification', '§', 'Test Title', 'Component', 
                  'Test by', 'Sample quantity', 'Test Timing', 'Start', 'End', 
                  'OK/NOK', 'Result', 'Remark', 'Customer Feedback']
        
        # 헤더 스타일
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 헤더 작성
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # 데이터 작성
        current_category = None
        row_num = 2
        
        for item in plan_data:
            # 카테고리 구분 행
            if item['_category'] != current_category:
                current_category = item['_category']
                ws.merge_cells(f'A{row_num}:O{row_num}')
                cell = ws.cell(row=row_num, column=1, value=current_category)
                cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                row_num += 1
            
            # 데이터 행
            for col, header in enumerate(headers, 1):
                value = item.get(header, '')
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
                # Functional Test 하이라이트
                if item['_is_functional']:
                    cell.fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
            
            row_num += 1
        
        # 열 너비 조정
        column_widths = [5, 12, 20, 8, 30, 15, 10, 12, 12, 10, 10, 8, 10, 15, 18]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width
        
        # 파일명 생성
        filename = f"Test_Plan_{user_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(filename)
        
        return filename

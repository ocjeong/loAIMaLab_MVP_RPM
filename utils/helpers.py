import pandas as pd
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def export_to_excel(plan_df, request_data):
    """계획서를 Excel 파일로 변환"""
    output = BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # 의뢰 정보 시트
        info_df = pd.DataFrame([{
            '발주처': request_data['client'],
            '프로젝트': request_data['project'],
            '작성일': request_data['created_date']
        }])
        info_df.to_excel(writer, sheet_name='의뢰정보', index=False)
        
        # 시험 계획 시트
        plan_df.to_excel(writer, sheet_name='시험계획', index=False)
        
        # 스타일 적용
        workbook = writer.book
        
        # 의뢰정보 시트 스타일
        info_sheet = workbook['의뢰정보']
        for cell in info_sheet[1]:
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # 시험계획 시트 스타일
        plan_sheet = workbook['시험계획']
        for cell in plan_sheet[1]:
            cell.font = Font(bold=True, size=11)
            cell.fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # 열 너비 자동 조정
        for sheet in workbook.worksheets:
            for column in sheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                sheet.column_dimensions[column_letter].width = adjusted_width
    
    output.seek(0)
    return output.getvalue()
